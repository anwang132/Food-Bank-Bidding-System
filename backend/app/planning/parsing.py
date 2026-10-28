"""Extract planning signals from uploaded files and typed notes.

Supported for analysis: CSV/TSV, Excel workbooks (.xlsx), Markdown, plain text and
JSON. Any other file is stored as-is and reported as "not analyzed".

What is extracted:

* a small table preview (spreadsheets, Markdown tables, JSON arrays of objects);
* **needs**: rows or lines that name an item and a quantity, e.g. a sheet with
  "Item" and "Cases needed" columns, or a note line ``- Infant formula: 30 cases``;
* categories and dietary terms mentioned anywhere in the text.

Untrusted input is bounded: workbooks are checked for decompressed size before
openpyxl reads them (it uses defusedxml when installed), and rows, cells, needs
and stored text are all capped.
"""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePosixPath
from typing import Any

from app.planning.catalog import (
    Category,
    classify,
    dietary_terms,
    mentioned_categories,
    parse_category,
)

MAX_ROWS = 5000
MAX_PREVIEW_ROWS = 20
MAX_COLUMNS = 30
MAX_CELL = 120
MAX_NEEDS = 200
MAX_TEXT = 50_000
MAX_SHEETS = 10
MAX_XLSX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_XLSX_ENTRIES = 2000

FORMATS = {
    ".csv": "csv",
    ".tsv": "csv",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".md": "markdown",
    ".markdown": "markdown",
    ".txt": "text",
    ".text": "text",
    ".json": "json",
}

_ITEM_HEADER = re.compile(r"item|product|food|description|commodity|name", re.I)
_QTY_HEADER = re.compile(r"qty|quantity|cases|units|amount|count|need|request", re.I)
_CATEGORY_HEADER = re.compile(r"category|type|group", re.I)
_NUMBER = re.compile(r"-?\d[\d,]*(?:\.\d+)?")
_UNITS = r"(?:cases?|boxes?|bags?|kits?|units?|sacks?|cans?|jars?|packs?|cartons?|lbs?|pounds?)"
# "- Infant formula: 30 cases", "1. Rice — 40", "Peanut butter = 12"
_LINE_ITEM_FIRST = re.compile(
    r"^\s*(?:[-*+•]\s+|\d+[.)]\s+)?(?P<item>[^\W\d_][\w ,'&/()-]{1,60}?)"
    r"\s*(?::|—|–|-|=)\s*(?P<qty>\d[\d,]*)\b"
)
# "30 cases of infant formula"
_LINE_QTY_FIRST = re.compile(
    rf"^\s*(?:[-*+•]\s+)?(?P<qty>\d[\d,]*)\s+{_UNITS}\s+(?:of\s+)?"
    r"(?P<item>[^\W\d_](?:[^\W\d_]|[ ,'-]){2,60})",
    re.I,
)


@dataclass
class Need:
    item: str
    quantity: int | None
    category: Category

    def as_dict(self) -> dict[str, Any]:
        return {"item": self.item, "quantity": self.quantity, "category": self.category}


@dataclass
class Analysis:
    format: str
    text: str = ""
    columns: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    row_count: int = 0
    sheets: list[str] = field(default_factory=list)
    needs: list[Need] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return {
            "format": self.format,
            "columns": self.columns,
            "rows": self.rows,
            "row_count": self.row_count,
            "sheets": self.sheets,
            "needs": [n.as_dict() for n in self.needs[:MAX_NEEDS]],
            "categories": mentioned_categories(self.text),
            "dietary": dietary_terms(self.text),
            "warnings": self.warnings,
        }


def detect_format(filename: str) -> str:
    return FORMATS.get(PurePosixPath(filename.lower()).suffix, "binary")


def analyze_file(filename: str, data: bytes) -> Analysis:
    fmt = detect_format(filename)
    if fmt == "xlsx":
        return _analyze_xlsx(data)
    if fmt == "binary":
        return Analysis(
            "binary", warnings=["Stored for reference; this file type is not analyzed."]
        )
    text, warning = _decode(data)
    if text is None:
        return Analysis(
            "binary", warnings=["The file is not readable text, so it was not analyzed."]
        )
    result = analyze_text(text, fmt, delimiter="\t" if filename.lower().endswith(".tsv") else None)
    if warning:
        result.warnings.insert(0, warning)
    return result


def analyze_text(text: str, fmt: str = "markdown", *, delimiter: str | None = None) -> Analysis:
    if fmt == "csv":
        return _analyze_csv(text, delimiter)
    if fmt == "json":
        return _analyze_json(text)
    return _analyze_prose(text, fmt)


# ---------------------------------------------------------------- formats


def _decode(data: bytes) -> tuple[str | None, str | None]:
    if b"\x00" in data[:8192]:
        return None, None
    try:
        return data.decode("utf-8-sig"), None
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace"), "Read with Windows-1252 encoding."


def _analyze_csv(text: str, delimiter: str | None) -> Analysis:
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = []
    for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter)):
        if i >= MAX_ROWS:
            break
        rows.append(row)
    result = Analysis("csv", text=text[:MAX_TEXT])
    _fill_table(result, rows)
    return result


def _analyze_xlsx(data: bytes) -> Analysis:
    result = Analysis("xlsx")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > MAX_XLSX_ENTRIES
                or sum(e.file_size for e in entries) > MAX_XLSX_UNCOMPRESSED
            ):
                result.warnings.append("The workbook is too large to analyze; it was stored only.")
                return result
    except zipfile.BadZipFile:
        result.warnings.append("This is not a valid .xlsx workbook; it was stored only.")
        return result

    from openpyxl import load_workbook  # deferred: only workbook uploads pay the import

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # openpyxl raises many types for malformed files
        result.warnings.append("The workbook could not be read; it was stored only.")
        return result
    try:
        lines: list[str] = []
        budget = MAX_ROWS
        # Preview the first sheet that lists needs, else the first sheet with data.
        preview: Analysis | None = None
        for sheet in workbook.worksheets[:MAX_SHEETS]:
            rows: list[list[str]] = []
            for values in sheet.iter_rows(values_only=True):
                if budget <= 0:
                    break
                budget -= 1
                row = [_cell(v) for v in values[:MAX_COLUMNS]]
                if any(row):
                    rows.append(row)
            if not rows:
                continue
            result.sheets.append(sheet.title)
            lines.extend("\t".join(r) for r in rows)
            table = Analysis("xlsx")
            _fill_table(table, rows)
            result.needs.extend(table.needs)
            if preview is None or (table.needs and not preview.needs):
                preview = table
        if preview is not None:
            result.columns, result.rows = preview.columns, preview.rows
            result.row_count = preview.row_count
        result.text = "\n".join(lines)[:MAX_TEXT]
    finally:
        workbook.close()
    if not result.sheets:
        result.warnings.append("The workbook has no data.")
    return result


def _analyze_json(text: str) -> Analysis:
    try:
        value = json.loads(text)
    except ValueError:
        result = _analyze_prose(text, "json")
        result.warnings.append("The file is not valid JSON; it was read as text.")
        return result
    records = value
    if isinstance(value, dict):
        records = next((v for v in value.values() if isinstance(v, list)), None)
    result = Analysis("json", text=json.dumps(value, indent=2, ensure_ascii=False)[:MAX_TEXT])
    if isinstance(records, list) and records and all(isinstance(r, dict) for r in records):
        columns: list[str] = []
        for record in records[:MAX_ROWS]:
            for key in record:
                if key not in columns and len(columns) < MAX_COLUMNS:
                    columns.append(str(key))
        rows = [columns] + [[_cell(r.get(c)) for c in columns] for r in records[:MAX_ROWS]]
        _fill_table(result, rows)
    return result


def _analyze_prose(text: str, fmt: str) -> Analysis:
    result = Analysis(fmt, text=text[:MAX_TEXT])
    table: list[list[str]] = []
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("|") and stripped.count("|") >= 2:
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if not all(set(c) <= set("-: ") for c in cells):  # skip |---|---| separators
                table.append(cells)
            continue
        if table:
            _take_table(result, table)
            table = []
        need = _need_from_line(stripped)
        if need is not None:
            result.needs.append(need)
    if table:
        _take_table(result, table)
    return result


def _take_table(result: Analysis, rows: list[list[str]]) -> None:
    table = Analysis(result.format)
    _fill_table(table, rows)
    result.needs.extend(table.needs)
    if not result.columns:
        result.columns, result.rows, result.row_count = table.columns, table.rows, table.row_count


# ---------------------------------------------------------------- tables and needs


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()[:MAX_CELL]


def _fill_table(result: Analysis, rows: list[list[str]]) -> None:
    rows = [[c.strip()[:MAX_CELL] for c in r[:MAX_COLUMNS]] for r in rows]
    rows = [r for r in rows if any(r)]
    header_at = next((i for i, r in enumerate(rows) if sum(1 for c in r if c) >= 2), None)
    if header_at is None:
        return
    header = rows[header_at]
    body = rows[header_at + 1 :]
    result.columns = header
    result.rows = body[:MAX_PREVIEW_ROWS]
    result.row_count = len(body)

    item_col = _find_column(header, _ITEM_HEADER)
    qty_col = _find_column(header, _QTY_HEADER, exclude=item_col)
    cat_col = _find_column(header, _CATEGORY_HEADER, exclude=item_col)
    if item_col is None:
        return
    for row in body:
        item = row[item_col] if item_col < len(row) else ""
        if not item or not re.search(r"[^\W\d_]", item):
            continue
        qty = _number(row[qty_col]) if qty_col is not None and qty_col < len(row) else None
        category = None
        if cat_col is not None and cat_col < len(row) and row[cat_col]:
            category = parse_category(row[cat_col])
        result.needs.append(Need(item[:80], qty, category or classify(item)))
        if len(result.needs) >= MAX_NEEDS:
            break


def _find_column(
    header: list[str], pattern: re.Pattern[str], exclude: int | None = None
) -> int | None:
    for i, name in enumerate(header):
        if i != exclude and pattern.search(name):
            return i
    return None


def _number(text: str) -> int | None:
    match = _NUMBER.search(text.replace(" ", ""))
    if not match:
        return None
    value = round(float(match.group().replace(",", "")))
    return value if 0 < value <= 1_000_000 else None


def _need_from_line(line: str) -> Need | None:
    for pattern in (_LINE_ITEM_FIRST, _LINE_QTY_FIRST):
        match = pattern.match(line)
        if match:
            item = match.group("item").strip(" ,-")
            qty = _number(match.group("qty"))
            if item and qty:
                return Need(item[:80], qty, classify(item))
    return None
